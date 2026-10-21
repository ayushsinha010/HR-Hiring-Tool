import os
import re

import docx
import fitz  # pymupdf


def extract_text(file_path):
    ext = os.path.splitext(file_path)[1].lower()
    if ext == ".pdf":
        return _extract_pdf(file_path)
    if ext == ".docx":
        return _extract_docx(file_path)
    if ext == ".txt":
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            return f.read()
    return ""


def _extract_pdf(file_path):
    text_parts = []
    with fitz.open(file_path) as doc:
        for page in doc:
            text_parts.append(page.get_text())
    return "\n".join(text_parts)


def _extract_docx(file_path):
    d = docx.Document(file_path)
    parts = [p.text for p in d.paragraphs if p.text.strip()]
    for table in d.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    parts.append(cell.text)
    return "\n".join(parts)


_NAME_NOISE_TOKENS = {"naukri", "resume", "cv", "updated", "latest", "final", "new"}
_BRACKETED = re.compile(r"[\[\(][^\]\)]*[\]\)]")  # strips "[6y 0m]" / "(6y 0m)" style annotations
_EMAIL_PATTERN = re.compile(r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}")


def extract_email(resume_text):
    """First email-looking string in the resume text, or None. Job-portal/LinkedIn export
    footers occasionally contain their own addresses, so this can occasionally pick up
    the wrong one - it's meant as a fast default, not a guarantee; HR can edit it."""
    match = _EMAIL_PATTERN.search(resume_text or "")
    return match.group(0) if match else None


def guess_candidate_name(filename, resume_text):
    stem = os.path.splitext(os.path.basename(filename))[0]
    stem = re.sub(r"^\d+[_\-\s]*", "", stem)  # leading numeric id, e.g. "108_"
    stem = _BRACKETED.sub("", stem)  # experience annotations, wherever they appear
    parts = [p for p in re.split(r"[_\-\s]+", stem) if p.strip()]
    parts = [p for p in parts if p.strip().lower() not in _NAME_NOISE_TOKENS]
    stem = " ".join(parts).strip()
    if stem:
        return stem
    first_line = next((l.strip() for l in resume_text.splitlines() if l.strip()), "")
    return first_line[:80] if first_line else os.path.basename(filename)
