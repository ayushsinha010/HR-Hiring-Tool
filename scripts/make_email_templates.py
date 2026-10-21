"""Regenerates the four editable candidate-email templates (.docx) in email_templates/ from the
built-in wording in emailer.py, using COMPANY_NAME / COMPANY_WEBSITE from the environment.

Run from the project folder:   python scripts/make_email_templates.py

The app reads these files on every send, so you can also just open them in Word and edit the text
(keep the "Subject Line" and "Email Body" headings and the {{ tokens }}). If a file is missing or
cannot be parsed, the app falls back to the same built-in wording.
"""
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import docx  # noqa: E402

import emailer  # noqa: E402
from config import COMPANY_NAME  # noqa: E402

DESCRIPTIONS = {
    "no_response": "Use this when you've tried calling a candidate about an opening but couldn't reach them - the Job Description is attached.",
    "jd_only": "Use this to send just the Job Description to a candidate - no interview details yet.",
    "link_only": "Use this to share only the interview date, time and meeting link - no Job Description attached.",
    "jd_and_link": "Use this to send the Job Description together with the interview date, time and meeting link.",
}


def build(kind):
    subject = emailer.EMAIL_SUBJECT_TEMPLATES[kind].format(role_name="{{ role_name }}", company=COMPANY_NAME)
    body = emailer._render_builtin_email(
        kind, "{{ candidate_name }}", "{{ role_name }}", "{{ sender_name }}", "{{ sender_email }}",
        date="{{ date }}", time="{{ time }}", teams_link="{{ teams_link }}",
    )
    document = docx.Document()
    document.core_properties.author = ""
    document.core_properties.title = f"Email Template - {emailer.EMAIL_TEMPLATE_LABELS[kind]}"
    document.add_heading(f"Email Template - {emailer.EMAIL_TEMPLATE_LABELS[kind]}", level=1)
    document.add_paragraph(DESCRIPTIONS[kind])
    document.add_heading("Subject Line", level=2)
    document.add_paragraph(subject)
    document.add_heading("Email Body", level=2)
    for paragraph in body.split("\n\n"):
        document.add_paragraph(paragraph)
    return document


def main():
    out_dir = emailer.EMAIL_TEMPLATE_DIR
    os.makedirs(out_dir, exist_ok=True)
    for kind in emailer.EMAIL_TEMPLATE_KINDS:
        path = emailer.email_template_path(kind)
        build(kind).save(path)
        print("wrote", path)


if __name__ == "__main__":
    main()
