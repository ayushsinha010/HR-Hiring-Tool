import json
import os
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

import db

HEADER_FILL = PatternFill(start_color="662984", end_color="662984", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)
TITLE_FONT = Font(bold=True, size=14, color="662984")

BUCKET_FILLS = {
    "green": PatternFill(start_color="F2FFF2", end_color="F2FFF2", fill_type="solid"),
    "yellow": PatternFill(start_color="FFF8E8", end_color="FFF8E8", fill_type="solid"),
    "red": PatternFill(start_color="FCE8E8", end_color="FCE8E8", fill_type="solid"),
    "unscored": PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid"),
}

CANDIDATE_COLUMNS = [
    ("Candidate ID", 12),
    ("Candidate Name", 26),
    ("JD Match Rating (out of 10)", 14),
    ("Screening Status", 14),
    ("Opinion on Profile (vs JD)", 45),
    ("Strengths", 45),
    ("Concerns", 45),
    ("Resume File", 30),
    ("Reached", 10),
    ("L1 Scheduled", 12),
    ("L1 Attended", 12),
    ("L1 Cleared", 12),
    ("L2 Scheduled", 12),
    ("L2 Attended", 12),
    ("L2 Cleared", 12),
    ("L3 Scheduled", 12),
    ("L3 Attended", 12),
    ("L3 Cleared", 12),
    ("Offer Made", 12),
    ("Joined", 10),
    ("Rejected", 10),
    ("On Hold", 10),
    ("Interview Video", 20),
    ("Interview Transcript", 20),
]


def _yn(value):
    return "Y" if value else "N"


def _awaiting_counts(candidates):
    return {
        "awaiting_l2": sum(1 for c in candidates if c["l1_cleared"] and not c["l2_scheduled"]),
        "awaiting_l3": sum(1 for c in candidates if c["l2_cleared"] and not c["l3_scheduled"]),
        "awaiting_offer": sum(1 for c in candidates if c["l3_cleared"] and not c["offer_made"]),
    }


def _write_ceo_dashboard(wb, position, kpis, awaiting):
    ws = wb.active
    ws.title = "CEO Dashboard"

    ws.append([f"HIRING TRACKER - CEO VIEW - {position['title']}"])
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=8)
    ws.cell(row=1, column=1).font = TITLE_FONT
    ws.append([f"Generated {datetime.now().strftime('%d-%b-%Y %H:%M')}"])
    ws.append([])

    headers = [
        "In Funnel", "Reached", "L1 Cleared", "L2 Cleared", "L3 Cleared",
        "Offers Made", "Joined", "Selection Rate",
    ]
    values = [
        kpis["total"], kpis["reached"], kpis["l1_cleared"], kpis["l2_cleared"],
        kpis["l3_cleared"], kpis["offer_made"], kpis["joined"], f"{kpis['selection_rate']}%",
    ]
    header_row = 4
    for col_idx, name in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=col_idx, value=name)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[cell.column_letter].width = 14
    for col_idx, value in enumerate(values, start=1):
        cell = ws.cell(row=header_row + 1, column=col_idx, value=value)
        cell.font = Font(bold=True, size=13)
        cell.alignment = Alignment(horizontal="center")

    awaiting_row = header_row + 3
    ws.cell(row=awaiting_row, column=1, value="Awaiting L2").font = Font(bold=True, color="797973")
    ws.cell(row=awaiting_row, column=2, value=awaiting["awaiting_l2"])
    ws.cell(row=awaiting_row + 1, column=1, value="Awaiting L3").font = Font(bold=True, color="797973")
    ws.cell(row=awaiting_row + 1, column=2, value=awaiting["awaiting_l3"])
    ws.cell(row=awaiting_row + 2, column=1, value="Awaiting Offer").font = Font(bold=True, color="797973")
    ws.cell(row=awaiting_row + 2, column=2, value=awaiting["awaiting_offer"])
    return ws


def _write_candidate_master(wb, candidates):
    ws = wb.create_sheet("Candidate Master")

    ws.append([f"CANDIDATE MASTER - one row per candidate"])
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(CANDIDATE_COLUMNS))
    ws.cell(row=1, column=1).font = TITLE_FONT
    ws.append([])

    header_row = 3
    for col_idx, (name, width) in enumerate(CANDIDATE_COLUMNS, start=1):
        cell = ws.cell(row=header_row, column=col_idx, value=name)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[cell.column_letter].width = width

    for row_idx, c in enumerate(candidates, start=header_row + 1):
        positives = ", ".join(json.loads(c["positives_json"])) if c["positives_json"] else ""
        negatives = ", ".join(json.loads(c["negatives_json"])) if c["negatives_json"] else ""
        values = [
            c["id"],
            c["candidate_name"],
            c["score"],
            c["bucket"],
            c["fit_summary"] or "",
            positives,
            negatives,
            c["resume_filename"],
            _yn(c["reached"]),
            _yn(c["l1_scheduled"]),
            _yn(c["l1_attended"]),
            _yn(c["l1_cleared"]),
            _yn(c["l2_scheduled"]),
            _yn(c["l2_attended"]),
            _yn(c["l2_cleared"]),
            _yn(c["l3_scheduled"]),
            _yn(c["l3_attended"]),
            _yn(c["l3_cleared"]),
            _yn(c["offer_made"]),
            _yn(c["joined"]),
            _yn(c["rejected"]),
            _yn(c["on_hold"]),
            c["interview_video_filename"] or "",
            c["interview_transcript_filename"] or "",
        ]
        for col_idx, value in enumerate(values, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.fill = BUCKET_FILLS.get(c["bucket"], BUCKET_FILLS["unscored"])
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    ws.freeze_panes = f"A{header_row + 1}"
    return ws


def build_workbook(position_id):
    """Two-sheet workbook mirroring the CEO Dashboard / Candidate Master shape of the
    original Sr. Python Master Hiring Tracker, populated from this app's own data."""
    position = db.get_position(position_id)
    candidates = db.list_candidates(position_id)
    kpis = db.funnel_for_position(position_id)
    awaiting = _awaiting_counts(candidates)

    wb = Workbook()
    _write_ceo_dashboard(wb, position, kpis, awaiting)
    _write_candidate_master(wb, candidates)
    return wb, position


def build_tracker(position_id):
    """Saves the workbook into this position's own data folder and returns the path."""
    wb, _position = build_workbook(position_id)
    out_dir = db.position_dir(position_id)
    out_path = os.path.join(out_dir, "tracker.xlsx")
    wb.save(out_path)
    return out_path


def export_to_desktop(position_id):
    """Saves a copy of the tracker workbook to this machine's Desktop and returns its path."""
    wb, position = build_workbook(position_id)
    desktop = Path.home() / "Desktop"
    desktop.mkdir(parents=True, exist_ok=True)
    safe_title = "".join(c for c in position["title"] if c.isalnum() or c in " -_").strip() or "tracker"
    filename = f"{safe_title} Tracker.xlsx"
    out_path = desktop / filename
    wb.save(out_path)
    return out_path
