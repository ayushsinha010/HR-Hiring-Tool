"""Optional template-based PDF reports (cover page, brand fonts, branded charts) - used only
when an external Markdown-to-Word template renderer (a folder holding scripts/build.py,
build_doc.py, inject_vml.py and build_toc.py, pointed to by the REPORT_TEMPLATE_DIR
environment variable), LibreOffice and Open Sans are all available on this machine. That
renderer is NOT part of this repository: with REPORT_TEMPLATE_DIR unset, is_available() is
False and the route uses report_generator.py (reportlab) instead.

Content is identical to report_generator.py's: the OVERALL report holds every table and graph
of the overall dashboard in dashboard order and ends with the last section "Hiring Timeline and
Buffer Time" (three by-position tables, an average-days-per-stage graph and two by-candidate
tables); a per-position report ends with that position's timeline rows. The numbers come from
report_data.py / timeline.py - the same builders the dashboard page uses.

How tables are kept from being cut off: the template renderer (build_doc.py) turns Markdown
tables into plain Word tables. Before the document is branded and converted, this module
re-opens it with python-docx and, for each table, sets a fixed layout with explicit column
widths (so long text wraps inside its cell instead of overflowing), a smaller font for wide
tables, "repeat as header row" on the first row and "do not split row across pages" on every
row. Charts are matplotlib PNGs (transparent, Open Sans) embedded as figures.
"""
import importlib.util
import math
import os
import shutil
import subprocess
import sys
import tempfile
from datetime import datetime

import db
import report_data
import timeline

TEMPLATE_RENDERER_DIR = os.environ.get("REPORT_TEMPLATE_DIR", "")
BUILD_SCRIPT = os.path.join(TEMPLATE_RENDERER_DIR, "scripts", "build.py")

FUNNEL_LABELS = ["Sourced", "Reached", "L1", "L2", "L3", "Offered", "Joined"]
FUNNEL_KEYS = ["total", "reached", "l1_cleared", "l2_cleared", "l3_cleared", "offer_made", "joined"]

CYAN, GREEN, ORANGE, PURPLE, RED, GREY, GRID = "#0096AF", "#96C800", "#D0811C", "#662984", "#B51729", "#555555", "#C9C9BE"
HEAT_COLORS = ["#EBEBDD", "#C6E9C6", "#8FD48F", "#96C800", "#5A7A00"]

PHASE_HEADERS = ["Sourced", "Reached", "Interview scheduled", "Inter-viewed", "Round cleared", "Offer made", "On hold", "Joined", "Rejected"]
PHASE_KEYS = list(db.STAGE_PHASES)


def is_available():
    return bool(TEMPLATE_RENDERER_DIR) and os.path.exists(BUILD_SCRIPT)


# ---------------------------------------------------------------------------
# charts (matplotlib, transparent so the cream page shows through)
# ---------------------------------------------------------------------------

def _setup_matplotlib():
    import matplotlib

    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib import font_manager

    font_dir = os.path.expandvars(r"%LOCALAPPDATA%\Microsoft\Windows\Fonts")
    for name in ("OpenSans-Regular.ttf", "OpenSans-Bold.ttf"):
        path = os.path.join(font_dir, name)
        if os.path.exists(path):
            font_manager.fontManager.addfont(path)
    plt.rcParams.update(
        {
            "font.family": "Open Sans",
            "text.color": GREY,
            "axes.labelcolor": GREY,
            "xtick.color": GREY,
            "ytick.color": GREY,
            "axes.edgecolor": GREY,
        }
    )
    return plt


def _style_axes(ax):
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    ax.tick_params(length=0)
    ax.set_axisbelow(True)


def _save(plt, fig, out_path):
    fig.tight_layout()
    fig.savefig(out_path, transparent=True)
    plt.close(fig)


def _bar_chart(out_path, labels, values, color):
    plt = _setup_matplotlib()
    fig, ax = plt.subplots(figsize=(7.2, 3.0), dpi=150)
    ax.bar(labels, values, color=color)
    _style_axes(ax)
    ax.yaxis.grid(True, color=GRID, linewidth=0.8)
    _save(plt, fig, out_path)


def _status_chart(out_path, counts):
    plt = _setup_matplotlib()
    labels = ["Active", "Closed", "Upcoming"]
    values = [counts["active"], counts["closed"], counts["upcoming"]]
    fig, ax = plt.subplots(figsize=(7.2, 3.0), dpi=150)
    ax.bar(labels, values, color=[GREEN, "#797973", ORANGE])
    _style_axes(ax)
    ax.yaxis.grid(True, color=GRID, linewidth=0.8)
    _save(plt, fig, out_path)


def _hbar_chart(out_path, labels, values, color, fmt="{:g}"):
    plt = _setup_matplotlib()
    n = max(len(values), 1)
    fig, ax = plt.subplots(figsize=(7.2, min(max(1.4 + 0.42 * n, 1.9), 6.5)), dpi=150)
    labels = [l if len(l) <= 42 else l[:41] + "..." for l in labels]
    bars = ax.barh(labels, values, color=color)
    ax.invert_yaxis()
    _style_axes(ax)
    ax.xaxis.grid(True, color=GRID, linewidth=0.8)
    peak = max(values) if values else 1
    for bar, v in zip(bars, values):
        ax.text(bar.get_width() + peak * 0.015, bar.get_y() + bar.get_height() / 2, fmt.format(v), va="center", fontsize=8)
    ax.set_xlim(0, peak * 1.12 if peak else 1)
    _save(plt, fig, out_path)


def _donut_chart(out_path, a_label, a_value, b_label, b_value):
    plt = _setup_matplotlib()
    fig, ax = plt.subplots(figsize=(4.6, 3.0), dpi=150)
    ax.pie([max(a_value, 0), max(b_value, 0)], colors=[GREEN, RED], startangle=90, counterclock=False,
           wedgeprops={"width": 0.38, "edgecolor": "none"})
    ax.legend([f"{a_label} ({a_value})", f"{b_label} ({b_value})"], loc="center left", bbox_to_anchor=(1.0, 0.5), frameon=False)
    ax.set_aspect("equal")
    _save(plt, fig, out_path)


def _heatmap_chart(out_path, heatmap):
    from matplotlib.patches import Rectangle

    plt = _setup_matplotlib()
    weeks = heatmap["weeks"]
    fig, ax = plt.subplots(figsize=(7.2, 1.7), dpi=150)
    for wi, week in enumerate(weeks):
        for di, day in enumerate(week):
            ax.add_patch(Rectangle((wi, -di), 0.86, 0.86, facecolor=HEAT_COLORS[day["level"]], edgecolor="none"))
    for m in heatmap["month_labels"]:
        ax.text(m["week"], 1.25, m["label"], fontsize=6, ha="left", va="bottom")
    for row, label in ((0, "Mon"), (2, "Wed"), (4, "Fri")):
        ax.text(-0.6, -row + 0.4, label, fontsize=5.5, ha="right", va="center")
    ax.set_xlim(-3.2, len(weeks) + 0.2)
    ax.set_ylim(-6.3, 2.6)
    ax.set_aspect("equal")
    ax.axis("off")
    _save(plt, fig, out_path)


def _position_cards_chart(out_path, funnel_rows):
    """One small funnel graph per position (the dashboard's position cards), two per row."""
    plt = _setup_matplotlib()
    n = len(funnel_rows)
    cols = 2
    rows = math.ceil(n / cols)
    fig, axes = plt.subplots(rows, cols, figsize=(7.2, 2.25 * rows), dpi=150, squeeze=False)
    for idx in range(rows * cols):
        ax = axes[idx // cols][idx % cols]
        if idx >= n:
            ax.axis("off")
            continue
        r = funnel_rows[idx]
        ax.bar(FUNNEL_LABELS, [r[k] for k in FUNNEL_KEYS], color=PURPLE)
        _style_axes(ax)
        ax.yaxis.grid(True, color=GRID, linewidth=0.6)
        ax.tick_params(axis="x", labelsize=6)
        ax.tick_params(axis="y", labelsize=6)
        days = f"{r['days_to_fill']}d to fill" if r["days_to_fill"] is not None else f"{r['days_open']}d open"
        title = r["title"] if len(r["title"]) <= 34 else r["title"][:33] + "..."
        ax.set_title(f"{title}\n{db.STATUS_LABELS[r['status']]} | {r['selection_rate']}% selection | {days}", fontsize=7, loc="left", color=PURPLE)
    _save(plt, fig, out_path)


# ---------------------------------------------------------------------------
# markdown helpers
# ---------------------------------------------------------------------------

def _cell(value):
    text = "-" if value is None else str(value)
    text = text.replace("|", "/").replace("*", "").replace("`", "").replace("\n", " ").strip()
    return text or "-"


class _Doc:
    """Collects the Markdown plus, for every table in document order, how it must be laid
    out (font size, relative column widths, header font) for the python-docx pass."""

    def __init__(self):
        self.parts = []
        self.table_specs = []

    def add(self, text):
        self.parts.append(text)

    def table(self, headers, rows, widths=None, font=8, header_font=None):
        header = "| " + " | ".join(_cell(h) for h in headers) + " |"
        sep = "|" + "|".join(["---"] * len(headers)) + "|"
        lines = [header, sep] + ["| " + " | ".join(_cell(c) for c in r) + " |" for r in rows]
        self.parts.append("\n".join(lines))
        self.table_specs.append({"font": font, "header_font": header_font or font, "widths": widths or [1] * len(headers)})

    def text(self):
        return "\n\n".join(self.parts) + "\n"


def _kpi_table(doc, pairs):
    doc.table([p[0] for p in pairs], [[p[1] for p in pairs]], font=10)


# ---------------------------------------------------------------------------
# sections
# ---------------------------------------------------------------------------

def _add_timeline_section(doc, tl, scratch_dir, heading_level="#", chart_name="avg_days_chart.png", show_position=True):
    tg = tl["targets"]
    doc.add(f"{heading_level} Hiring Timeline and Buffer Time")
    doc.add(
        "When the JD went up, when each CV came in, where every CV is now and how long things have been waiting, in calendar "
        "days. A dash (-) means that date was never recorded; **est.** means no stage date existed, so the CV-added date stands in."
    )
    doc.add(
        f"**Over target when:** in one stage more than {tg['stage_age_days']}d; feedback pending more than "
        f"{tg['feedback_pending_days']}d; no decision after feedback more than {tg['decision_pending_days']}d; position open more "
        f"than {tg['position_open_days']}d. **{tl['totals']['candidates_over_target']} of {tl['totals']['candidates']} candidate(s) are over target.**"
    )

    sub = heading_level + "#"
    doc.add(f"{sub} By position")
    doc.add("JD and CVs, and how long each position has been open.")
    rows = []
    for p in tl["positions"]:
        open_val = p["days_to_fill"] if p["filled"] and p["days_to_fill"] is not None else p["open_days"]
        rows.append([p["title"], timeline.fmt_date(p["jd_uploaded_on"]), timeline.fmt_date(p["first_cv_on"]),
                     timeline.fmt_days(p["days_jd_to_first_cv"]), p["cvs_received"],
                     f"{timeline.fmt_days(open_val)}{' (filled)' if p['filled'] else ''}", p["items_over_target"]])
    doc.table(["Position", "JD uploaded on", "First CV on", "JD to 1st CV (d)", "CVs received", "Position open (d)", "Items over target"],
              rows or [["No positions yet."] + ["-"] * 6], widths=[3, 2.1, 2.1, 1.4, 1.4, 1.6, 1.4], font=8)

    doc.add("Candidates in each stage.")
    rows = [[p["title"]] + [p["stage_counts"][ph] for ph in PHASE_KEYS] for p in tl["positions"]]
    doc.table(["Position"] + PHASE_HEADERS, rows or [["No positions yet."] + ["-"] * 9],
              widths=[2.6] + [1.15] * 9, font=7.5, header_font=6.5)

    doc.add("Interviews, feedback and average days per stage.")
    rows = []
    for p in tl["positions"]:
        avg = "; ".join(f"{k} {v}d" for k, v in p["avg_days_per_stage"].items()) or "-"
        rows.append([p["title"], p["interviews_scheduled"], p["interviews_done"], p["feedback_saved"], p["feedback_pending"], avg])
    doc.table(["Position", "Interviews scheduled", "Interviews done", "Feedback saved", "Feedback pending", "Average days per stage"],
              rows or [["No positions yet."] + ["-"] * 5], widths=[2.4, 1.3, 1.3, 1.3, 1.3, 4.2], font=8)

    if tl["avg_days_per_stage"]:
        chart_path = os.path.join(scratch_dir, chart_name)
        _hbar_chart(chart_path, [f"{r['phase']} (n={r['count']})" for r in tl["avg_days_per_stage"]],
                    [r["avg_days"] for r in tl["avg_days_per_stage"]], CYAN, fmt="{:.1f}")
        doc.add(f"{sub} Average days per stage")
        doc.add("How long candidates have spent in each stage (completed stays plus the stage they are in now); estimated dates are left out.")
        doc.add(f"![Figure. Average days per stage.]({chart_name})")

    doc.add(f"{sub} By candidate (most delayed first)")
    doc.add("Where each CV is now and what is waiting.")
    rows = []
    for c in tl["candidates"]:
        est = " est." if c["entered_is_estimate"] else ""
        if c["over_target"]:
            flag = "OVER TARGET: " + "; ".join(c["over_target_reasons"])
        elif c["phase"] in ("Joined", "Rejected", "On Hold"):
            flag = "-"
        else:
            flag = "On track"
        full = [c["name"], c["position_title"], timeline.fmt_date(c["cv_fed_on"]), c["stage_label"],
                timeline.fmt_date(c["entered_stage_on"]) + est,
                timeline.fmt_days(c["days_in_stage"]) + (est if c["days_in_stage"] is not None else ""),
                c["next_interview"], c["feedback_status"], flag]
        rows.append(full if show_position else full[:1] + full[2:])
    headers = ["Candidate", "Position", "CV fed on", "Current stage", "Entered stage on", "Days in stage", "Next interview", "Feedback", "Over target"]
    widths = [1.9, 1.5, 1.1, 1.5, 1.1, 0.7, 1.6, 1.0, 2.3]
    if not show_position:   # a position's own report hides the redundant Position column
        headers, widths = headers[:1] + headers[2:], widths[:1] + widths[2:]
    doc.table(headers, rows or [["No candidates yet."] + ["-"] * (len(headers) - 1)], widths=widths, font=6.5, header_font=6)

    doc.add("Elapsed days per candidate.")
    rows = [[c["name"], timeline.fmt_days(c["days_since_cv"]), timeline.fmt_days(c["days_cv_to_schedule"]),
             timeline.fmt_days(c["days_interview_to_feedback"]), timeline.fmt_days(c["days_feedback_to_decision"])]
            for c in tl["candidates"]]
    doc.table(["Candidate", "Days since CV fed", "CV to interview scheduled (d)", "Interview to feedback (d)", "Feedback to decision (d)"],
              rows or [["No candidates yet."] + ["-"] * 4], widths=[3, 1.3, 1.6, 1.5, 1.5], font=8)


# ---------------------------------------------------------------------------
# docx post-processing + the build pipeline
# ---------------------------------------------------------------------------

def _style_tables(docx_path, table_specs):
    """Fixed layout with explicit widths (long text wraps inside its cell), a smaller font for wide
    tables, repeat-header-row on the first row, rows that never split across pages."""
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, Twips

    doc = Document(docx_path)
    tables = doc.tables
    if len(tables) != len(table_specs):
        raise RuntimeError(f"table count mismatch: document has {len(tables)}, expected {len(table_specs)}")
    total = 9360  # 6.5 inch of text width, in twips
    for tbl, spec in zip(tables, table_specs):
        rel = spec["widths"]
        widths = [max(int(total * r / sum(rel)), 300) for r in rel]
        tblPr = tbl._tbl.tblPr
        look = tblPr.find(qn("w:tblLook"))
        for tag, attrs in (("w:tblW", {"w:w": str(sum(widths)), "w:type": "dxa"}),):
            el = tblPr.find(qn(tag))
            if el is None:
                el = OxmlElement(tag)
                (look.addprevious(el) if look is not None else tblPr.append(el))
            for k, v in attrs.items():
                el.set(qn(k), v)
        layout = OxmlElement("w:tblLayout")
        layout.set(qn("w:type"), "fixed")
        mar = OxmlElement("w:tblCellMar")
        for side, w in (("top", 25), ("left", 45), ("bottom", 25), ("right", 45)):
            m = OxmlElement(f"w:{side}")
            m.set(qn("w:w"), str(w))
            m.set(qn("w:type"), "dxa")
            mar.append(m)
        for el in (layout, mar):
            (look.addprevious(el) if look is not None else tblPr.append(el))
        grid = tbl._tbl.tblGrid
        for gc, w in zip(grid.findall(qn("w:gridCol")), widths):
            gc.set(qn("w:w"), str(w))
        for ri, row in enumerate(tbl.rows):
            trPr = row._tr.get_or_add_trPr()
            trPr.append(OxmlElement("w:cantSplit"))
            if ri == 0:
                trPr.append(OxmlElement("w:tblHeader"))
            size = Pt(spec["header_font"] if ri == 0 else spec["font"])
            for ci, cell in enumerate(row.cells):
                cell.width = Twips(widths[ci])
                for p in cell.paragraphs:
                    p.paragraph_format.space_after = Pt(0)
                    p.paragraph_format.space_before = Pt(0)
                    for run in p.runs:
                        run.font.size = size
    doc.save(docx_path)


def _load_renderer_build_module():
    spec = importlib.util.spec_from_file_location("template_renderer_build", BUILD_SCRIPT)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def _run_build(scratch_dir, md_filename, deliverable_name, table_specs=None):
    """Markdown -> branded PDF through the template renderer's own scripts. With table_specs, the
    build stages are run one by one so the tables can be laid out (see _style_tables) between
    the Word build and the branding/TOC/PDF stages; without it the renderer's one-command
    wrapper runs unchanged."""
    if not table_specs:
        result = subprocess.run(
            [sys.executable, BUILD_SCRIPT, md_filename, deliverable_name, "--no-preview"],
            cwd=scratch_dir, capture_output=True, encoding="utf-8", errors="replace", timeout=240,
        )
        pdf_path = os.path.join(scratch_dir, "output", deliverable_name, deliverable_name + ".pdf")
        if not os.path.exists(pdf_path):
            raise RuntimeError("Template build failed.\n" + (result.stdout or "") + "\n" + (result.stderr or ""))
        return pdf_path

    build = _load_renderer_build_module()
    if not build.ensure_on_path("soffice", build.SOFFICE_DIRS):
        raise RuntimeError("soffice (LibreOffice) not found.")
    build.ensure_on_path("pdftotext", build.POPPLER_DIRS)
    build.ensure_on_path("pdftoppm", build.POPPLER_DIRS)
    scripts = os.path.join(TEMPLATE_RENDERER_DIR, "scripts")
    raw = os.path.join(scratch_dir, "_raw.docx")
    brand = os.path.join(scratch_dir, "_branded.docx")
    docx = os.path.join(scratch_dir, deliverable_name + ".docx")
    pdf = os.path.join(scratch_dir, deliverable_name + ".pdf")

    def stage(name, argv):
        r = subprocess.run(argv, cwd=scratch_dir, capture_output=True, encoding="utf-8", errors="replace", timeout=240)
        if r.returncode != 0:
            raise RuntimeError(f"Template build stage '{name}' failed.\n{r.stdout}\n{r.stderr}")

    stage("build_doc", [sys.executable, os.path.join(scripts, "build_doc.py"), os.path.join(scratch_dir, md_filename), raw])
    _style_tables(raw, table_specs)
    stage("inject_vml", [sys.executable, os.path.join(scripts, "inject_vml.py"), raw, brand])
    stage("build_toc", [sys.executable, os.path.join(scripts, "build_toc.py"), brand, docx])
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", scratch_dir, docx],
                   cwd=scratch_dir, capture_output=True, timeout=240)
    if not os.path.exists(pdf):
        raise RuntimeError("Template build failed: soffice produced no PDF.")
    return pdf


def _front_matter(title, subtitle):
    return f"---\ntitle: {title}\nsubtitle: {subtitle}\nyear: {datetime.now().year}\n---\n\n"


# ---------------------------------------------------------------------------
# the two reports
# ---------------------------------------------------------------------------

def generate_position_report(position_id):
    position = db.get_position(position_id)
    candidates = db.list_candidates(position_id)
    kpis = db.funnel_for_position(position_id)
    tl = timeline.build_timeline([position_id])

    scratch_dir = tempfile.mkdtemp(prefix="branded_report_")
    try:
        doc = _Doc()
        chart_path = os.path.join(scratch_dir, "funnel_chart.png")
        _bar_chart(chart_path, FUNNEL_LABELS, [kpis[k] for k in FUNNEL_KEYS], CYAN)

        status_label = db.STATUS_LABELS[kpis["status"]]
        doc.add("# Summary")
        doc.add(f"{position['title']} is currently **{status_label}**, with {kpis['total']} candidate(s) in the funnel and a {kpis['selection_rate']}% selection rate.")
        _kpi_table(doc, [
            ("In Funnel", kpis["total"]), ("Reached", kpis["reached"]), ("L1 Cleared", kpis["l1_cleared"]),
            ("L2 Cleared", kpis["l2_cleared"]), ("L3 Cleared", kpis["l3_cleared"]), ("Offers", kpis["offer_made"]),
            ("Joined", kpis["joined"]), ("Selection Rate", f"{kpis['selection_rate']}%"),
        ])
        doc.add("# Hiring Funnel")
        doc.add("The chart below tracks candidates from sourcing through to joining.")
        doc.add(f"![Figure 1. Candidate funnel for {position['title']}.](funnel_chart.png)")

        doc.add("# Candidates")
        doc.add("Every candidate currently in this position's pipeline, highest score first; resumes with no score yet come last.")
        rows = [[c["candidate_name"] or c["resume_filename"],
                 f"{c['score']:.1f}" if c["score"] is not None else "Not scored yet", c["bucket"], db.current_stage_label(c)]
                for c in candidates]
        doc.table(["Candidate", "Score", "Screening", "Current Stage"], rows or [["No candidates yet.", "-", "-", "-"]],
                  widths=[3.4, 1.3, 1.3, 2], font=9)

        _add_timeline_section(doc, tl, scratch_dir, show_position=False)
        md = _front_matter(f"Hiring Report \u2014 {position['title']}", f"Generated {datetime.now().strftime('%d %b %Y')}") + doc.text()
        with open(os.path.join(scratch_dir, "report.md"), "w", encoding="utf-8") as f:
            f.write(md)

        safe_name = "".join(c for c in position["title"] if c.isalnum() or c in " -_").strip() or "Hiring Report"
        pdf_path = _run_build(scratch_dir, "report.md", safe_name, doc.table_specs)
        with open(pdf_path, "rb") as f:
            return f.read(), position["title"]
    finally:
        shutil.rmtree(scratch_dir, ignore_errors=True)


def generate_overall_report(position_ids=None):
    """The Overall Hiring Report: every table and graph of the overall dashboard, then the Hiring
    Timeline and Buffer Time section last. Returns the PDF bytes."""
    data = report_data.overall_dashboard_data(position_ids=position_ids)
    counts, overall, funnel_rows = data["counts"], data["overall"], data["funnel_rows"]
    stage_groups, status_fields = data["stage_groups"], data["status_fields"]

    scratch_dir = tempfile.mkdtemp(prefix="branded_report_")
    try:
        doc = _Doc()
        _status_chart(os.path.join(scratch_dir, "status_chart.png"), counts)
        _bar_chart(os.path.join(scratch_dir, "funnel_chart.png"), FUNNEL_LABELS, [overall[k] for k in FUNNEL_KEYS], CYAN)

        doc.add("# Summary")
        doc.add(f"Across {counts['total_positions']} position(s), {overall['total']} candidate(s) are in the funnel with an overall selection rate of {overall['selection_rate']}%.")
        _kpi_table(doc, [("Total Positions", counts["total_positions"]), ("Active Hiring", counts["active"]),
                         ("Closed Hiring", counts["closed"]), ("Upcoming Hiring", counts["upcoming"])])
        compact = [("Candidates in Funnel", overall["total"]), ("Joined", overall["joined"]),
                   ("Overall Selection Rate", f"{overall['selection_rate']}%")]
        if data["avg_days_to_hire"] is not None:
            compact.append(("Avg Days to Hire", f"{data['avg_days_to_hire']}d"))
        _kpi_table(doc, compact)

        doc.add("# Positions by Status")
        doc.add("![Figure 1. Positions grouped by hiring status.](status_chart.png)")
        doc.add("# Overall Hiring Funnel")
        doc.add("![Figure 2. Candidate funnel across all positions.](funnel_chart.png)")

        heat = data["activity_heatmap"]
        _heatmap_chart(os.path.join(scratch_dir, "heatmap.png"), heat)
        doc.add("# Activity Heatmap")
        doc.add(f"{heat['total']} action(s) in the last year. Each square is one day; darker means more hiring activity (stage changes, notes, emails, new candidates).")
        doc.add("![Figure 3. Activity heatmap, last 12 months.](heatmap.png)")

        fig = 3
        filled = [r for r in funnel_rows if r["days_to_fill"] is not None]
        open_rows = [r for r in funnel_rows if r["days_to_fill"] is None]
        if filled:
            fig += 1
            _hbar_chart(os.path.join(scratch_dir, "time_to_fill.png"), [r["title"] for r in filled], [r["days_to_fill"] for r in filled], CYAN)
            doc.add("# Time to Fill")
            doc.add("Days from opening to filled.")
            doc.add(f"![Figure {fig}. Time to fill, by position.](time_to_fill.png)")
        if open_rows:
            fig += 1
            _hbar_chart(os.path.join(scratch_dir, "age_open.png"), [r["title"] for r in open_rows], [r["days_open"] for r in open_rows], ORANGE)
            doc.add("# Age of Open Positions")
            doc.add("Days each position has been open.")
            doc.add(f"![Figure {fig}. Age of open positions.](age_open.png)")
        hire = data["days_to_hire_by_position"]
        if hire:
            fig += 1
            _hbar_chart(os.path.join(scratch_dir, "time_to_hire.png"), [r["title"] for r in hire], [r["avg_days"] for r in hire], PURPLE, fmt="{:.1f}")
            doc.add("# Time to Hire by Role")
            doc.add("Average days from sourced to joined.")
            doc.add(f"![Figure {fig}. Time to hire by role.](time_to_hire.png)")
        offer = data["offer_acceptance"]
        if offer and offer.get("rate") is not None:
            fig += 1
            _donut_chart(os.path.join(scratch_dir, "offer_acceptance.png"), "Accepted", offer["accepted"], "Declined", offer["declined"])
            doc.add("# Offer Acceptance Rate")
            doc.add(f"{offer['accepted']} accepted, {offer['declined']} declined ({offer['rate']}% acceptance).")
            doc.add(f"![Figure {fig}. Offer acceptance rate.](offer_acceptance.png)")

        if funnel_rows:
            fig += 1
            _position_cards_chart(os.path.join(scratch_dir, "position_cards.png"), funnel_rows)
            doc.add("# Positions")
            doc.add("One funnel graph per position, with its status, selection rate and days open or to fill.")
            doc.add(f"![Figure {fig}. Hiring funnel by position.](position_cards.png)")

        doc.add("# Full Funnel by Position")
        doc.add("Every pipeline stage per position; the last row adds them all up. Sch = Scheduled, Att = Attended, Clr = Cleared, Rej = Rejected, Hold = On Hold.")
        short = {"Scheduled": "Sch", "Attended": "Att", "Cleared": "Clr"}
        headers = ["Position", "Days", "In funnel", "Reached"]
        for group, items in stage_groups:
            if group.startswith("L"):
                headers += [f"{group[:2]} {short[label]}" for _f, label, _g in items]
        headers += ["Offer", "Joined", "Rej.", "Hold"]
        fields = ["reached"] + [f for g, items in stage_groups if g.startswith("L") for f, _l, _x in items] + ["offer_made", "joined", "rejected", "on_hold"]
        rows = []
        for r in funnel_rows:
            days = f"{r['days_to_fill']}d (filled)" if r["days_to_fill"] is not None else f"{r['days_open']}d"
            rows.append([r["title"], days, r["total"]] + [r[f] for f in fields])
        rows.append(["TOTAL - all positions", "-", overall["total"]] + [overall[f] for f in fields])
        doc.table(headers, rows, widths=[2.4, 1.4, 1.1, 1.35] + [1.0] * (len(fields) - 1), font=6.5, header_font=5.5)

        _add_timeline_section(doc, data["timeline"], scratch_dir)   # the LAST section

        md = _front_matter("Overall Hiring Report", f"Generated {datetime.now().strftime('%d %b %Y')}") + doc.text()
        with open(os.path.join(scratch_dir, "report.md"), "w", encoding="utf-8") as f:
            f.write(md)
        pdf_path = _run_build(scratch_dir, "report.md", "Overall Hiring Report", doc.table_specs)
        with open(pdf_path, "rb") as f:
            return f.read()
    finally:
        shutil.rmtree(scratch_dir, ignore_errors=True)
